"""proposal.json → 공모전 제안서 .docx

사용: python build_docx.py <proposal.json> <out.docx>
의존성: python-docx
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "맑은 고딕"
HEAD_FILL = "E4DFEC"   # 연보라 — 공모전 양식 머리행 색
TOTAL_FILL = "F2F2F2"
ACCENT = RGBColor(0x3B, 0x2A, 0x6B)


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_rich(par, text, size=10.5, bold=False, color=None):
    """`**굵게**` 만 해석한다."""
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if chunk:
            set_font(par.add_run(chunk), size, bold or i % 2 == 1, color)


def shade(cell, fill):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def par_border(par, side="bottom", sz=12, color="3B2A6B"):
    p_pr = par._element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{side}")
    for k, v in {"val": "single", "sz": str(sz), "space": "4", "color": color}.items():
        el.set(qn(f"w:{k}"), v)
    bdr.append(el)
    p_pr.append(bdr)


def spacing(par, before=0, after=4, line=1.15):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def write_cell(cell, text, bold=False, size=9.5, align=None):
    cell.text = ""
    par = cell.paragraphs[0]
    spacing(par, 1, 1, 1.1)
    if align is not None:
        par.alignment = align
    add_rich(par, str(text), size, bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def is_short(text):
    return len(str(text)) <= 12


def table(doc, b):
    cols = b["columns"]
    rows = list(b["rows"])
    total = b.get("total")
    t = doc.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cols):
        write_cell(t.rows[0].cells[i], c, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(t.rows[0].cells[i], HEAD_FILL)
    for r in rows + ([total] if total else []):
        cells = t.add_row().cells
        last = r is total
        for i, v in enumerate(r):
            align = WD_ALIGN_PARAGRAPH.CENTER if is_short(v) else None
            write_cell(cells[i], v, bold=last, align=align)
            if last:
                shade(cells[i], TOTAL_FILL)
    if b.get("merge_first_col"):
        start = 1
        for i in range(2, len(t.rows) - (1 if total else 0)):
            if str(rows[i - 1][0]).strip() == "":
                t.cell(start, 0).merge(t.cell(i, 0))
                # merge 는 빈 문단을 이어붙이므로 첫 문단만 남긴다
                cell = t.cell(start, 0)
                for p in cell.paragraphs[1:]:
                    p._element.getparent().remove(p._element)
            else:
                start = i
    if b.get("widths"):
        for row in t.rows:
            for i, w in enumerate(b["widths"]):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def block(doc, b, base):
    kind = b["type"]
    if kind == "paragraph":
        p = doc.add_paragraph()
        spacing(p, 2, 4, 1.3)
        add_rich(p, b["text"])
    elif kind == "subheading":
        p = doc.add_paragraph()
        spacing(p, 8, 3)
        add_rich(p, b["text"], 11, True)
    elif kind == "bullets":
        for item in b["items"]:
            p = doc.add_paragraph()
            spacing(p, 0, 2, 1.25)
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.first_line_indent = Cm(-0.35)
            add_rich(p, "- " + item, 10)
    elif kind == "note":
        p = doc.add_paragraph()
        spacing(p, 2, 6)
        add_rich(p, "※ " + b["text"].lstrip("※ "), 9, color=RGBColor(0x55, 0x55, 0x55))
    elif kind == "callout":
        t = doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        c = t.rows[0].cells[0]
        write_cell(c, b["text"], size=10.5)
        shade(c, TOTAL_FILL)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    elif kind == "table":
        table(doc, b)
    elif kind == "figure":
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(p, 4, 0)
        p.add_run().add_picture(str(base / b["path"]), width=Cm(b.get("width_cm", 16)))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(cap, 2, 8)
        add_rich(cap, b["caption"], 9, color=RGBColor(0x55, 0x55, 0x55))
    else:
        raise ValueError(f"알 수 없는 블록: {kind}")


def build(src, out):
    src = Path(src)
    data = json.loads(src.read_text(encoding="utf-8"))
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(1.8)

    if data.get("contest"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing(p, 0, 2)
        add_rich(p, data["contest"], 9.5, color=RGBColor(0x66, 0x66, 0x66))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(p, 4, 2)
    add_rich(p, data["title"], 17, True, ACCENT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    spacing(p, 0, 8)
    par_border(p)
    add_rich(p, f"( 제안서명: {data['name']} )", 12, True)

    block(doc, {"type": "callout", "text": data["summary"]}, src.parent)

    for s in data["sections"]:
        h = doc.add_paragraph()
        spacing(h, 12, 6)
        par_border(h, sz=6, color="9A8FB8")
        add_rich(h, s["heading"], 13, True, ACCENT)
        for b in s["blocks"]:
            block(doc, b, src.parent)

    if data.get("author"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        spacing(p, 12, 0)
        add_rich(p, data["author"], 9.5)

    doc.save(out)
    print(out)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("사용: python build_docx.py <proposal.json> <out.docx>")
    build(sys.argv[1], sys.argv[2])
